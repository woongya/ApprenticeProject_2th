from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Manufacturing_AI_Motor_ML_Project.pptx"


def add_title(slide, title: str, subtitle: str = ""):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.2), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 36, 66)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.0), Inches(0.35))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = RGBColor(80, 80, 80)


def add_bullets(slide, items: list[str], x=0.8, y=1.45, w=11.6, h=4.9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.space_after = Pt(8)


def add_table(slide, rows: list[list[str]], x=0.45, y=1.35, w=12.45, h=4.8):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    widths = [2.4, 5.4, 4.6]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.text_frame.paragraphs[0].font.size = Pt(13 if r else 15)
            cell.text_frame.paragraphs[0].font.bold = r == 0 or c == 0
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(219, 230, 242)
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return table


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "제조 AI 머신러닝 프로젝트", "MOVELOG 기반 모터 부하/관성/제어 지표 예측")
    add_bullets(slide, [
        "문제: Motor Driver의 1ms MOVELOG CSV를 이용해 I, L, KI를 예측한다.",
        "입력 X: target, resistance, accel, speed, pwm 조건과 speed/current/pwm waveform feature",
        "출력 Y: inertia(I), load(L), KI label",
        "활용: 모터 상태 판단과 parameter recommendation의 기초 모델 구축",
    ])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Project 1", "제조 AI 머신러닝 프로젝트 수행 절차")
    add_table(slide, [
        ["단계", "주요 내용", "단계별 미션"],
        ["1. 데이터셋 획득 및 문제 정의", "MOVELOG CSV 사용\n입력변수 X와 출력변수 Y 정의\n회귀 문제로 시작", "프로젝트 대상 데이터와 예측 문제를 명확히 설정"],
        ["2. 데이터 분할 및 비교", "데이터 수가 작으면 Leave-One-Out CV\n충분하면 Train/Test 또는 Train/Validation/Test", "검증 데이터 유무에 따른 결과 차이 분석"],
        ["3. 하이퍼파라미터 조정", "RandomForest Grid Search\nn_estimators, max_depth, min_samples_leaf 조정", "튜닝 전/후 성능 비교"],
        ["4. 데이터 스케일링 및 이유 설명", "Tree model은 scale 영향이 작음\nScaler는 train에만 fit해야 함", "Data Leakage 방지 설명"],
        ["5. 최종 성능 평가", "MAE, RMSE, R2\n+/-1 accuracy\nfeature importance 분석", "모델 일반화 성능을 객관적으로 검증"],
    ], h=5.2)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Dataset", "MOVELOG CSV 구조")
    add_bullets(slide, [
        "CSV metadata: run_id, repeat, target, resistance, accel, max_speed, max_pwm, current_limit, inertia, load, ki",
        "Raw columns: slot, index, time_ms, pos_delta, control_speed, pwm_out, current, flags",
        "CSV 파일 하나를 하나의 학습 sample로 사용한다.",
        "파일명과 metadata에 포함된 I/L/KI를 label로 사용한다.",
    ])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Feature Extraction", "Raw waveform을 학습 가능한 table로 변환")
    add_bullets(slide, [
        "속도 feature: speed_peak, speed_avg, speed_std, speed_p90, rise_time",
        "전류 feature: current_peak, current_avg, current_std, current_ripple, current_p90",
        "PWM feature: pwm_peak, pwm_avg, speed_per_pwm_gain",
        "동작 feature: duration_ms, pos_total, energy_proxy, moving_ratio, fault_ratio",
        "출력 파일: outputs/dataset_features.csv",
    ])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Modeling", "I, L, KI 예측 모델")
    add_bullets(slide, [
        "Baseline 모델: MultiOutputRegressor(RandomForestRegressor)",
        "데이터 수가 20개 이상이면 train/test split과 GridSearchCV 사용",
        "데이터 수가 작으면 Leave-One-Out Cross Validation 사용",
        "평가 지표: MAE, RMSE, R2, +/-1 accuracy",
        "feature_importance.csv로 어떤 feature가 중요한지 확인",
    ])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Experiment Plan", "데이터 수집 및 검증 계획")
    add_bullets(slide, [
        "Batch Test로 Resistance, Accel, Max Speed를 변경하며 MOVELOG 수집",
        "Repeat은 전체 parameter sweep 이후 다시 반복하여 환경 편향을 줄임",
        "품질 체크: sample_count, current_peak, speed_peak, pos_total, fault flag",
        "초기 데이터로 baseline을 만들고, 데이터가 늘어날수록 성능 변화를 비교",
    ])

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Expected Result", "최종 산출물")
    add_bullets(slide, [
        "MOVELOG raw CSV dataset",
        "dataset_features.csv feature table",
        "I/L/KI 예측 모델 model_i_l_ki.joblib",
        "evaluation_report.txt 성능 평가 결과",
        "향후 확장: 예측된 I/L/KI 기반 parameter 추천",
    ])

    prs.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
